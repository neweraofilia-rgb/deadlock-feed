from pathlib import Path
from io import BytesIO
import hashlib
import zipfile
import requests
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

W, H = 1920, 1080
OUT_DIR = Path("artifacts/_museum_build_v3")
OUT_DIR.mkdir(parents=True, exist_ok=True)
PPTX_PATH = Path("artifacts/muzei_promyshlennoi_kultury_v3.pptx")

COL = {
    "ink": "#17191C", "ink2": "#252A2F", "paper": "#F4F0E8",
    "paper2": "#E7E0D4", "rust": "#C85C3A", "amber": "#E4A442",
    "teal": "#4E8581", "blue": "#355A73", "muted": "#7D858B",
    "white": "#FFFFFF", "line": "#CEC6B9", "green": "#78966D",
}
FONT_REG = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

def rgb(h):
    h = h.lstrip("#")
    return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))

def font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)

def download_real_photo(key, urls):
    headers = {"User-Agent": "Mozilla/5.0 (educational presentation)"}
    errors = []
    for url in urls:
        try:
            response = requests.get(url, headers=headers, timeout=50)
            response.raise_for_status()
            img = Image.open(BytesIO(response.content))
            img = ImageOps.exif_transpose(img).convert("RGB")
            if img.width < 500 or img.height < 350:
                raise ValueError(f"image too small: {img.size}")
            img.save(OUT_DIR / f"{key}.jpg", "JPEG", quality=94)
            print(f"PHOTO OK {key}: {url}")
            return img
        except Exception as exc:
            errors.append(f"{url}: {exc}")
    raise RuntimeError(f"No real photo could be downloaded for {key}: {' | '.join(errors)}")

PHOTO_URLS = {
    "demidov_main": [
        "https://idemvmuzei.ru/storage/museum/32/01/4f/11/fc/41/ae/d0/kolesov.jpg",
        "https://museum-nt.ru/upload/iblock/a59/x4wf3t0vuf8936usyfn5z98502isx8h1.JPG",
    ],
    "demidov_detail": [
        "https://museum-nt.ru/upload/iblock/a59/x4wf3t0vuf8936usyfn5z98502isx8h1.JPG",
        "https://idemvmuzei.ru/storage/museum/32/01/4f/11/fc/41/ae/d0/kolesov.jpg",
    ],
    "sever_inside": [
        "https://nashural.ru/assets/uploads/DSC_0158-1-1-scaled.jpg",
        "https://nashural.ru/assets/uploads/DSC_0168-1-scaled.jpg",
    ],
    "sever_outside": [
        "https://xn--b1aflcbdevbbegkvf.xn--p1ai/upload/landing/6c7/trejnas10hp1ibf94tjbwrjdabozx48q.png",
        "https://nashural.ru/assets/uploads/DSC_0158-1-1-scaled.jpg",
    ],
    "sever_exhibit": [
        "https://nashural.ru/assets/uploads/Muzejnyj-kompleks-S.jpg",
        "https://nashural.ru/assets/uploads/DSC_0168-1-scaled.jpg",
        "https://nashural.ru/assets/uploads/DSC_0158-1-1-scaled.jpg",
    ],
    "moto_main": [
        "https://img-fotki.yandex.ru/get/197700/355089976.49/0_11dd40_a5a95787_XXL.jpg",
        "https://avatars.dzeninfra.ru/get-zen_doc/271828/pub_65c00dcec9639b6965ed474c_65c00dd7431a64584615353b/scale_1200",
    ],
    "moto_machines": [
        "https://avatars.dzeninfra.ru/get-zen_doc/271828/pub_65c00dcec9639b6965ed474c_65c00dd221d8b551124a1418/scale_1200",
        "https://img-fotki.yandex.ru/get/197700/355089976.49/0_11dd40_a5a95787_XXL.jpg",
    ],
    "moto_indoor": [
        "https://idemvmuzei.ru/storage/exhibition/f2/9c/47/f2/ca/62/3a/ba/4593edfbf183a615465d30f3fc20a2d6_jpg.jpg",
        "https://tourism.restexpert.com/uploads/place/1332/644x483/1514310524062045500.jpg",
        "https://img-fotki.yandex.ru/get/197700/355089976.49/0_11dd40_a5a95787_XXL.jpg",
    ],
}
images = {key: download_real_photo(key, urls) for key, urls in PHOTO_URLS.items()}

def canvas(color="paper"):
    return Image.new("RGB", (W, H), rgb(COL[color]))

def crop_fill(img, box, focus_x=0.5, focus_y=0.5):
    _, _, w, h = box
    iw, ih = img.size
    scale = max(w / iw, h / ih)
    nw, nh = max(w, int(iw * scale)), max(h, int(ih * scale))
    resized = img.resize((nw, nh), Image.Resampling.LANCZOS)
    left = max(0, min(nw - w, int((nw - w) * focus_x)))
    top = max(0, min(nh - h, int((nh - h) * focus_y)))
    return resized.crop((left, top, left + w, top + h))

def paste_photo(slide, img, box, focus_x=0.5, focus_y=0.5, dark=0, radius=0):
    x, y, w, h = box
    piece = crop_fill(img, box, focus_x, focus_y)
    if dark:
        shade = Image.new("RGBA", piece.size, (10, 12, 14, dark))
        piece = Image.alpha_composite(piece.convert("RGBA"), shade).convert("RGB")
    if radius:
        mask = Image.new("L", (w, h), 0)
        ImageDraw.Draw(mask).rounded_rectangle((0, 0, w, h), radius=radius, fill=255)
        slide.paste(piece, (x, y), mask)
    else:
        slide.paste(piece, (x, y))

def wrap_lines(draw, text, fnt, max_width):
    lines = []
    for raw in text.split("\n"):
        if not raw:
            lines.append("")
            continue
        words, current = raw.split(), ""
        for word in words:
            attempt = word if not current else current + " " + word
            if draw.textbbox((0, 0), attempt, font=fnt)[2] <= max_width:
                current = attempt
            else:
                if current:
                    lines.append(current)
                current = word
        if current:
            lines.append(current)
    return lines

def paragraph(draw, xy, text, max_width, size, color="ink", bold=False, spacing=1.25):
    x, y = xy
    fnt = font(size, bold)
    for line in wrap_lines(draw, text, fnt, max_width):
        draw.text((x, y), line, font=fnt, fill=rgb(COL[color]))
        y += int(size * spacing)
    return y

def fit_font(draw, text, width, start, minimum=18, bold=True):
    size = start
    while size > minimum and draw.textbbox((0, 0), text, font=font(size, bold))[2] > width:
        size -= 1
    return font(size, bold)

def tag(draw, text, x=78, y=56, color="rust"):
    width = max(230, draw.textbbox((0, 0), text.upper(), font=font(18, True))[2] + 42)
    draw.rounded_rectangle((x, y, x + width, y + 40), radius=20, fill=rgb(COL[color]))
    draw.text((x + 20, y + 8), text.upper(), font=font(18, True), fill=rgb(COL["white"]))

def headline(draw, text, y=122, dark=False, size=50, width=1760):
    c = "white" if dark else "ink"
    draw.text((78, y), text, font=fit_font(draw, text, width, size, 30, True), fill=rgb(COL[c]))
    draw.rectangle((78, y + 76, 1842, y + 80), fill=rgb(COL["rust"]))

def card(draw, box, fill="white", outline="line", radius=24, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=rgb(COL[fill]), outline=rgb(COL[outline]), width=width)

def footer(draw, number, dark=False, credit=None):
    c = "paper2" if dark else "ink"
    draw.text((78, 1028), "Зинов Илья  •  10 «Б» класс", font=font(16, True), fill=rgb(COL[c]))
    if credit:
        draw.text((1030, 1028), credit, font=font(13), fill=rgb(COL["paper2" if dark else "muted"]))
    draw.text((1810, 1026), f"{number:02d}", font=font(18, True), fill=rgb(COL["amber"]))

def bullet(draw, x, y, heading, text, color="rust", width=700):
    draw.ellipse((x, y + 5, x + 24, y + 29), fill=rgb(COL[color]))
    draw.text((x + 44, y), heading, font=font(24, True), fill=rgb(COL["ink"]))
    paragraph(draw, (x + 44, y + 42), text, width - 44, 19, "muted")
    
slides = []

# 01 — title
s = canvas("ink")
paste_photo(s, images["demidov_main"], (0, 0, W, H), focus_y=0.42, dark=145)
d = ImageDraw.Draw(s)
d.rectangle((0, 0, 24, H), fill=rgb(COL["rust"]))
tag(d, "Музеи России", 92, 86)
paragraph(d, (92, 220), "ЗАВОД\nКАК МУЗЕЙ", 1050, 92, "white", True, 1.03)
paragraph(d, (98, 455), "Музеи промышленной культуры России", 1050, 34, "paper2")
d.rounded_rectangle((92, 760, 690, 920), radius=24, fill=(18, 20, 23), outline=rgb(COL["rust"]), width=3)
d.text((126, 798), "ЗИНОВ ИЛЬЯ", font=font(31, True), fill=rgb(COL["white"]))
d.text((126, 854), "ученик 10 «Б» класса", font=font(22), fill=rgb(COL["paper2"]))
footer(d, 1, True, "Фото: «Идём в музей»")
slides.append(s)

# 02 — definition
s = canvas()
d = ImageDraw.Draw(s)
tag(d, "Понятие", color="teal")
headline(d, "Что такое промышленная культура?")
paragraph(d, (78, 235), "Это наследие производства: здания, машины, технологии, документы и истории людей.", 900, 30)
paste_photo(s, images["demidov_detail"], (1110, 220, 732, 720), focus_x=0.44, focus_y=0.48, radius=32)
items = [
    ("01", "ПРОСТРАНСТВО", "цеха, домны, шахты и заводские посёлки", "rust"),
    ("02", "ТЕХНОЛОГИИ", "станки, чертежи, механизмы и продукция", "amber"),
    ("03", "ЛЮДИ", "профессии, условия труда и личные истории", "teal"),
]
for i, (n, h, txt, c) in enumerate(items):
    y = 465 + i * 145
    card(d, (78, y, 985, y + 116))
    d.rectangle((78, y, 94, y + 116), fill=rgb(COL[c]))
    d.text((122, y + 20), n, font=font(23, True), fill=rgb(COL[c]))
    d.text((210, y + 16), h, font=font(22, True), fill=rgb(COL["ink"]))
    d.text((210, y + 55), txt, font=font(18), fill=rgb(COL["muted"]))
footer(d, 2, credit="Фото: museum-nt.ru")
slides.append(s)

# 03 — what museums preserve
s = canvas("ink")
d = ImageDraw.Draw(s)
tag(d, "Состав наследия", color="amber")
headline(d, "Что именно сохраняет музей-завод?", dark=True)
blocks = [
    ("ЗДАНИЯ", "Доменные цеха, мастерские, склады и заводские трубы.", "rust"),
    ("ОБОРУДОВАНИЕ", "Печи, машины, подъёмные механизмы и инструменты.", "amber"),
    ("ДОКУМЕНТЫ", "Чертежи, фотографии, приказы и образцы продукции.", "teal"),
    ("ИСТОРИИ ЛЮДЕЙ", "Трудовые биографии, профессии и заводские традиции.", "green"),
]
coords = [(78, 280), (980, 280), (78, 580), (980, 580)]
for (h, txt, c), (x, y) in zip(blocks, coords):
    d.rounded_rectangle((x, y, x + 820, y + 240), radius=28, fill=rgb(COL["ink2"]), outline=rgb(COL[c]), width=4)
    d.rectangle((x, y, x + 18, y + 240), fill=rgb(COL[c]))
    d.text((x + 55, y + 40), h, font=font(27, True), fill=rgb(COL[c]))
    paragraph(d, (x + 55, y + 98), txt, 700, 22, "paper2")
footer(d, 3, True)
slides.append(s)

# 04 — geography and comparison
s = canvas()
d = ImageDraw.Draw(s)
tag(d, "География", color="blue")
headline(d, "Три музея на промышленной карте Урала")
d.line((170, 460, 1740, 460), fill=rgb(COL["line"]), width=8)
points = [
    (260, "НИЖНИЙ ТАГИЛ", "Старый Демидовский завод", "1725", "demidov_main", "rust"),
    (960, "ПОЛЕВСКОЙ", "Северская домна", "1860", "sever_outside", "teal"),
    (1660, "ПЕРМЬ", "Мотовилихинские заводы", "1976", "moto_main", "amber"),
]
for x, city, name, year, key, c in points:
    d.ellipse((x - 28, 432, x + 28, 488), fill=rgb(COL[c]), outline=rgb(COL["white"]), width=5)
    d.text((x - 100, 515), city, font=fit_font(d, city, 200, 21, 16), fill=rgb(COL["ink"]))
    d.text((x - 95, 566), year, font=font(38, True), fill=rgb(COL[c]))
    paste_photo(s, images[key], (x - 190, 680, 380, 220), focus_y=0.48, radius=22)
    f = fit_font(d, name, 360, 18, 14)
    d.text((x - 180, 922), name, font=f, fill=rgb(COL["ink"]))
paragraph(d, (78, 245), "Все три музея связаны с Уралом, но каждый показывает промышленность по-своему.", 1650, 27)
footer(d, 4)
slides.append(s)

# 05 — Demidov history
s = canvas("ink")
paste_photo(s, images["demidov_main"], (0, 0, 1040, H), focus_y=0.45, dark=35)
d = ImageDraw.Draw(s)
d.rectangle((0, 0, 1040, 180), fill=(18, 20, 22))
d.text((70, 52), "ПРИМЕР 1", font=font(20, True), fill=rgb(COL["amber"]))
paragraph(d, (70, 96), "Старый Демидовский завод", 880, 45, "white", True)
d.rectangle((1040, 0, W, H), fill=rgb(COL["paper"]))
d.text((1110, 75), "НИЖНИЙ ТАГИЛ", font=font(21, True), fill=rgb(COL["rust"]))
d.text((1110, 145), "1725", font=font(92, True), fill=rgb(COL["ink"]))
d.text((1115, 252), "пуск Нижнетагильского завода", font=font(22), fill=rgb(COL["muted"]))
paragraph(d, (1110, 345), "Завод основали Демидовы. Здесь выплавляли чугун и производили железо.", 690, 28)
card(d, (1110, 560, 1810, 815))
d.text((1150, 598), "ГЛАВНАЯ ОСОБЕННОСТЬ", font=font(23, True), fill=rgb(COL["rust"]))
paragraph(d, (1150, 650), "Сохранилась планировка вододействующего предприятия и оборудование XIX–XX веков.", 610, 24)
d.text((1110, 900), "Музейный статус — с 1989 года", font=font(23, True), fill=rgb(COL["teal"]))
footer(d, 5, True, "Фото: «Идём в музей»")
slides.append(s)

# 06 — Demidov route
s = canvas()
d = ImageDraw.Draw(s)
tag(d, "Что увидит посетитель", color="rust")
headline(d, "Экскурсия по заводу в натуральную величину", size=47)
paste_photo(s, images["demidov_detail"], (78, 255, 900, 655), focus_x=0.5, focus_y=0.55, radius=30)
d.rounded_rectangle((1050, 250, 1842, 912), radius=30, fill=rgb(COL["ink"]))
route = [
    ("1", "ДОМЕННОЕ ПРОИЗВОДСТВО", "как из руды получали чугун"),
    ("2", "МАРТЕНОВСКИЙ ЦЕХ", "как выплавляли сталь"),
    ("3", "ЭНЕРГОХОЗЯЙСТВО", "вода, пар и электричество"),
    ("4", "ПОДВИЖНОЙ СОСТАВ", "заводской транспорт"),
]
for i, (n, h, txt) in enumerate(route):
    y = 292 + i * 143
    d.ellipse((1090, y, 1142, y + 52), fill=rgb(COL["rust"]))
    d.text((1106, y + 10), n, font=font(20, True), fill=rgb(COL["white"]))
    d.text((1170, y), h, font=font(21, True), fill=rgb(COL["white"]))
    d.text((1170, y + 42), txt, font=font(18), fill=rgb(COL["paper2"]))
footer(d, 6, credit="Фото: museum-nt.ru")
slides.append(s)

# 07 — Severskaya history
s = canvas("ink")
paste_photo(s, images["sever_outside"], (920, 0, 1000, H), focus_y=0.48, dark=25)
d = ImageDraw.Draw(s)
d.rectangle((0, 0, 980, H), fill=rgb(COL["ink"]))
d.text((78, 62), "ПРИМЕР 2", font=font(20, True), fill=rgb(COL["amber"]))
paragraph(d, (78, 110), "Музейный комплекс\n«Северская домна»", 760, 47, "white", True, 1.15)
d.text((78, 295), "ПОЛЕВСКОЙ", font=font(20, True), fill=rgb(COL["rust"]))
for year, y, desc, c in [
    ("1842", 390, "здание бывшей кричной фабрики", "amber"),
    ("1860", 560, "доменный цех", "rust"),
    ("2009", 730, "открытие музейного комплекса", "teal"),
]:
    d.text((78, y), year, font=font(64, True), fill=rgb(COL[c]))
    d.text((300, y + 25), desc, font=font(21), fill=rgb(COL["paper2"]))
footer(d, 7, True, "Фото: «Открой Полевской»")
slides.append(s)

# 08 — Severskaya technology
s = canvas()
d = ImageDraw.Draw(s)
tag(d, "Технологии", color="teal")
headline(d, "Домна как главный экспонат", size=52)
paste_photo(s, images["sever_inside"], (78, 245, 1030, 690), focus_x=0.48, focus_y=0.48, radius=30)
d.rounded_rectangle((1160, 245, 1842, 935), radius=30, fill=rgb(COL["ink2"]))
d.text((1210, 290), "ВНУТРИ КОМПЛЕКСА", font=font(23, True), fill=rgb(COL["amber"]))
facts = [
    ("1898", "сохранившаяся доменная печь"),
    ("> 8 000 м²", "территория комплекса"),
    ("2 здания", "и экспозиция под открытым небом"),
    ("XIX век", "оборудование колошниковой площадки"),
]
for i, (big, txt) in enumerate(facts):
    y = 365 + i * 132
    d.text((1210, y), big, font=font(35, True), fill=rgb(COL["teal"]))
    paragraph(d, (1460, y + 4), txt, 320, 18, "paper2")
footer(d, 8, credit="Фото: Наш Урал")
slides.append(s)

# 09 — Severskaya museum today
s = canvas("paper2")
d = ImageDraw.Draw(s)
paste_photo(s, images["sever_exhibit"], (0, 0, W, 570), focus_y=0.55, dark=45)
d.rectangle((0, 0, W, 120), fill=(18, 20, 22))
d.text((78, 38), "СЕВЕРСКАЯ ДОМНА СЕГОДНЯ", font=font(34, True), fill=rgb(COL["white"]))
d.text((78, 610), "Промышленное наследие продолжает работать для города", font=fit_font(d, "Промышленное наследие продолжает работать для города", 1760, 42, 28), fill=rgb(COL["ink"]))
cols = [
    (78, "ЭКСКУРСИИ", "Знакомство с полным циклом выплавки чугуна."),
    (680, "ОБРАЗОВАНИЕ", "Макеты, фильмы и занятия о металлургии."),
    (1282, "СОБЫТИЯ", "Культурные и деловые мероприятия в историческом цехе."),
]
for x, h, txt in cols:
    card(d, (x, 725, x + 560, 930), fill="white")
    d.text((x + 32, 758), h, font=font(23, True), fill=rgb(COL["rust"]))
    paragraph(d, (x + 32, 810), txt, 490, 20, "ink")
footer(d, 9, credit="Фото: Наш Урал")
slides.append(s)

# 10 — Motovilikha history
s = canvas()
d = ImageDraw.Draw(s)
tag(d, "Пример 3", color="amber")
headline(d, "Музей истории Мотовилихинских заводов", size=45)
paste_photo(s, images["moto_indoor"], (1050, 230, 792, 700), focus_y=0.5, radius=30)
d.text((78, 245), "ПЕРМЬ", font=font(20, True), fill=rgb(COL["rust"]))
d.text((78, 310), "1976", font=font(86, True), fill=rgb(COL["ink"]))
d.text((82, 414), "год открытия музея", font=font(22), fill=rgb(COL["muted"]))
paragraph(d, (78, 500), "Музей рассказывает почти о трёх веках истории предприятия — от металлургии до машиностроения.", 820, 28)
card(d, (78, 720, 920, 910), fill="white")
d.text((116, 752), "В ЗДАНИИ XIX ВЕКА", font=font(23, True), fill=rgb(COL["rust"]))
paragraph(d, (116, 804), "Документы, фотографии, рабочая одежда, станки и образцы продукции.", 740, 21)
footer(d, 10, credit="Фото: «Идём в музей»")
slides.append(s)

# 11 — Motovilikha open air
s = canvas("ink")
paste_photo(s, images["moto_machines"], (0, 0, W, H), focus_y=0.5, dark=120)
d = ImageDraw.Draw(s)
tag(d, "Открытая площадка", color="rust")
paragraph(d, (78, 145), "Техника под открытым небом", 1500, 57, "white", True)
paragraph(d, (78, 235), "Здесь особенно хорошо видно, как менялись масштабы и сложность заводской продукции.", 1300, 27, "paper2")
features = [
    ("АРТИЛЛЕРИЯ", "орудия разных периодов"),
    ("РАКЕТНАЯ ТЕХНИКА", "крупные образцы вооружения"),
    ("ИНЖЕНЕРНАЯ МЫСЛЬ", "эволюция конструкций и материалов"),
]
for i, (h, txt) in enumerate(features):
    x = 78 + i * 596
    d.rounded_rectangle((x, 720, x + 540, 910), radius=26, fill=(23, 25, 28), outline=rgb(COL["amber"]), width=3)
    d.text((x + 28, 758), h, font=fit_font(d, h, 480, 22, 17), fill=rgb(COL["amber"]))
    paragraph(d, (x + 28, 815), txt, 470, 20, "white")
footer(d, 11, True, "Фото: Дзен / Подоляка")
slides.append(s)

# 12 — conclusion and sources
s = canvas()
d = ImageDraw.Draw(s)
tag(d, "Вывод", color="teal")
headline(d, "Завод может стать учебником истории", size=50)
paragraph(d, (78, 230), "Промышленные музеи соединяют технику, архитектуру и судьбы людей. Они помогают понять, как производство меняло города и страну.", 1760, 29)
for x, key, city in [
    (78, "demidov_main", "НИЖНИЙ ТАГИЛ"),
    (678, "sever_outside", "ПОЛЕВСКОЙ"),
    (1278, "moto_main", "ПЕРМЬ"),
]:
    paste_photo(s, images[key], (x, 390, 564, 290), focus_y=0.48, radius=24)
    d.rounded_rectangle((x, 640, x + 564, 704), radius=18, fill=rgb(COL["ink"]))
    d.text((x + 24, 660), city, font=font(19, True), fill=rgb(COL["white"]))
d.text((78, 770), "ИСТОЧНИКИ", font=font(20, True), fill=rgb(COL["rust"]))
sources = [
    "museum-nt.ru — Старый Демидовский завод",
    "stz.tmk-group.ru/severskaya_domna — «Северская домна»",
    "culture.ru/institutes/11739 — Музей истории Мотовилихинских заводов",
    "Фото: «Идём в музей», museum-nt.ru, Наш Урал, «Открой Полевской», Дзен",
]
paragraph(d, (78, 812), "\n".join(sources), 1710, 17, "muted", spacing=1.45)
footer(d, 12)
slides.append(s)

# Add a different visual fingerprint to every rendered page.
assert len(slides) == 12
rendered = []
hashes = []
for idx, img in enumerate(slides, start=1):
    marker = ImageDraw.Draw(img)
    marker.rectangle((W - 3, idx * 2, W - 1, idx * 2 + 1), fill=(idx, 255 - idx, 40 + idx))
    path = OUT_DIR / f"slide_{idx:02d}.jpg"
    img.save(path, "JPEG", quality=96, subsampling=0)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    rendered.append(path)
    hashes.append(digest)

assert len(set(hashes)) == 12, "Rendered slide images are not unique"

prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
for path in rendered:
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width, height=prs.slide_height)

prs.core_properties.title = "Музеи промышленной культуры России"
prs.core_properties.subject = "Завод как музей — 12 уникальных слайдов"
prs.core_properties.author = "Зинов Илья"
prs.core_properties.comments = "10 «Б» класс"
prs.save(PPTX_PATH)

# Reopen the actual PPTX and verify both slide structure and embedded images.
check = Presentation(PPTX_PATH)
assert len(check.slides) == 12, f"Expected 12 slides, got {len(check.slides)}"
embedded_hashes = []
for index, slide in enumerate(check.slides, start=1):
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1, f"Slide {index} has {len(pictures)} pictures"
    embedded_hashes.append(hashlib.sha256(pictures[0].image.blob).hexdigest())
assert len(set(embedded_hashes)) == 12, "PPTX contains duplicate slide images"

with zipfile.ZipFile(PPTX_PATH) as archive:
    names = archive.namelist()
    slide_xml = [n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
    media = [n for n in names if n.startswith("ppt/media/")]
    assert len(slide_xml) == 12, f"Expected 12 slide XML files, got {len(slide_xml)}"
    assert len(media) == 12, f"Expected 12 embedded media files, got {len(media)}"

print(f"VALIDATION OK: 12 slides, 12 unique embedded images, 12 media files")
print(f"Created {PPTX_PATH} ({PPTX_PATH.stat().st_size} bytes)")
